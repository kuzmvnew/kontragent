"""Fixture-only Mintrans TED registry ingestion and fact projection."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timezone
from hashlib import sha256
import json
import os
from pathlib import Path
import re
import shutil
from typing import Mapping
from zipfile import is_zipfile

from openpyxl import load_workbook
from sqlalchemy import delete, select

from app.models.company import Company
from app.models.mintrans_ted import (
    MINTRANS_TED_FACT_CODE,
    MintransTedEntry,
    MintransTedQuarantineRow,
    MintransTedRawArtifact,
    TransportForwardingRegistryListing,
)
from app.models.source import DataSet, DatasetPublication
from app.services.bulk_update_service import (
    BulkUpdatePipeline,
    IntegrityError,
    ParseValidationError,
    ParsedSnapshot,
    validate_artifact,
)
from app.services.mintrans_ted_registry_service import DATASET_CODE


MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MANIFEST_SCHEMA_VERSION = 1


class MintransTedSchemaError(ParseValidationError):
    code = "invalid_schema"


class ArtifactImmutabilityError(IntegrityError):
    code = "artifact_immutability_violation"


@dataclass(frozen=True)
class StagedFixtureArtifact:
    path: Path
    manifest_path: Path
    checksum: str
    size_bytes: int
    manifest: dict
    source_metadata: dict


@dataclass(frozen=True)
class MintransTedParsedArtifact:
    entries: tuple[dict, ...]
    quarantine: tuple[dict, ...]
    coverage: dict


@dataclass(frozen=True)
class IdentityMatch:
    company_id: int | None
    state: str
    method: str | None


def _canonical_json(value: object) -> str:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )


def _json_safe(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Mapping):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(item) for item in value]
    return str(value)


def _aware_utc(value: datetime) -> datetime:
    if value.tzinfo is None or value.utcoffset() is None:
        raise ValueError("source_as_of должен содержать часовой пояс")
    return value.astimezone(timezone.utc)


def _write_once(path: Path, payload: bytes) -> None:
    try:
        with path.open("xb") as stream:
            stream.write(payload)
            stream.flush()
            os.fsync(stream.fileno())
    except FileExistsError:
        if path.read_bytes() != payload:
            raise ArtifactImmutabilityError(
                f"Content-addressed artifact member already differs: {path.name}"
            )


def stage_fixture_artifact(
    source_path: str | Path,
    *,
    artifact_store: str | Path,
    source_as_of: datetime,
    source_metadata: Mapping[str, object] | None = None,
    expected_sha256: str | None = None,
) -> StagedFixtureArtifact:
    """Copy a local XLSX once into a checksum-addressed immutable directory."""

    source = Path(source_path)
    if source.suffix.lower() != ".xlsx":
        raise IntegrityError("Mintrans TED fixture должен иметь формат XLSX")
    integrity = validate_artifact(source, expected_sha256=expected_sha256)
    if not is_zipfile(source):
        raise IntegrityError("Файл с расширением XLSX не является XLSX-контейнером")

    source_as_of = _aware_utc(source_as_of)
    metadata = {
        "source_code": DATASET_CODE,
        "source_owner": "Минтранс России",
        "ingestion_mode": "fixture",
        "live_ingestion": False,
        **{
            str(key): _json_safe(value)
            for key, value in (source_metadata or {}).items()
        },
    }
    # Callers cannot relabel this bounded implementation as live ingestion.
    metadata["source_code"] = DATASET_CODE
    metadata["ingestion_mode"] = "fixture"
    metadata["live_ingestion"] = False

    manifest = {
        "schema_version": MANIFEST_SCHEMA_VERSION,
        "dataset_code": DATASET_CODE,
        "artifact_kind": "raw_fixture",
        "data_format": "xlsx",
        "media_type": MEDIA_TYPE,
        "sha256": integrity.checksum,
        "size_bytes": integrity.size_bytes,
        "original_file_name": source.name,
        "source_as_of": source_as_of.isoformat(),
        "source_metadata": metadata,
    }

    artifact_dir = Path(artifact_store) / integrity.checksum
    artifact_dir.mkdir(parents=True, exist_ok=True)
    target = artifact_dir / "artifact.xlsx"
    manifest_path = artifact_dir / "manifest.json"

    if target.exists():
        existing = validate_artifact(target, expected_sha256=integrity.checksum)
        if existing.size_bytes != integrity.size_bytes:
            raise ArtifactImmutabilityError("Размер ранее сохранённого RAW artifact изменён")
    else:
        try:
            with source.open("rb") as input_stream, target.open("xb") as output_stream:
                shutil.copyfileobj(input_stream, output_stream, length=1024 * 1024)
                output_stream.flush()
                os.fsync(output_stream.fileno())
        except Exception:
            target.unlink(missing_ok=True)
            raise
        validate_artifact(target, expected_sha256=integrity.checksum)

    manifest_bytes = (_canonical_json(manifest) + "\n").encode("utf-8")
    _write_once(manifest_path, manifest_bytes)
    return StagedFixtureArtifact(
        path=target,
        manifest_path=manifest_path,
        checksum=integrity.checksum,
        size_bytes=integrity.size_bytes,
        manifest=manifest,
        source_metadata=metadata,
    )


def _header_key(value: object) -> str:
    text = str(value or "").strip().lower().replace("ё", "е")
    return re.sub(r"\s+", " ", text)


_COLUMN_ALIASES = {
    "inn": {"инн", "inn"},
    "ogrn": {"огрн/огрнип", "огрн / огрнип", "огрн", "ogrn", "ogrn/ogrnip"},
    "registry_number": {
        "регистрационный номер",
        "номер реестровой записи",
        "реестровый номер",
        "registry number",
        "registry_number",
    },
    "included_at": {
        "дата включения",
        "дата включения в реестр",
        "included_at",
        "included at",
    },
}


def _resolve_columns(header: tuple[object, ...]) -> dict[str, int]:
    resolved: dict[str, int] = {}
    for index, value in enumerate(header):
        key = _header_key(value)
        for canonical, aliases in _COLUMN_ALIASES.items():
            if key in aliases:
                if canonical in resolved:
                    raise MintransTedSchemaError(
                        f"Семантическая колонка {canonical} указана более одного раза"
                    )
                resolved[canonical] = index
    missing = sorted(set(_COLUMN_ALIASES).difference(resolved))
    if missing:
        raise MintransTedSchemaError(
            "Отсутствуют обязательные колонки: " + ", ".join(missing)
        )
    return resolved


def _identifier_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    text = str(value).strip().replace("\u00a0", "")
    if text.startswith("'"):
        text = text[1:]
    return text or None


def is_valid_inn(value: str) -> bool:
    if not re.fullmatch(r"\d{10}|\d{12}", value):
        return False
    digits = [int(char) for char in value]
    if len(digits) == 10:
        weights = (2, 4, 10, 3, 5, 9, 4, 6, 8)
        return sum(a * b for a, b in zip(digits, weights)) % 11 % 10 == digits[9]
    weights_11 = (7, 2, 4, 10, 3, 5, 9, 4, 6, 8)
    weights_12 = (3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8)
    check_11 = sum(a * b for a, b in zip(digits, weights_11)) % 11 % 10
    check_12 = sum(a * b for a, b in zip(digits, weights_12)) % 11 % 10
    return check_11 == digits[10] and check_12 == digits[11]


def is_valid_ogrn(value: str) -> bool:
    if not re.fullmatch(r"\d{13}|\d{15}", value):
        return False
    if len(value) == 13:
        return int(value[:12]) % 11 % 10 == int(value[12])
    return int(value[:14]) % 13 % 10 == int(value[14])


def _registry_number(value: object) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _included_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    return None


def _hash_payload(value: object) -> str:
    return sha256(_canonical_json(value).encode("utf-8")).hexdigest()


def _normalize_row(raw: dict, *, row_number: int) -> tuple[dict | None, dict | None]:
    inn = _identifier_text(raw.get("inn"))
    ogrn = _identifier_text(raw.get("ogrn"))
    registry_number = _registry_number(raw.get("registry_number"))
    included_at = _included_date(raw.get("included_at"))
    reasons = []
    if inn is not None and not is_valid_inn(inn):
        reasons.append("invalid_inn")
    if ogrn is not None and not is_valid_ogrn(ogrn):
        reasons.append("invalid_ogrn")
    if inn is None and ogrn is None:
        reasons.append("missing_identity")
    if registry_number is None:
        reasons.append("missing_registry_number")
    elif len(registry_number) > 200:
        reasons.append("invalid_registry_number")
    if included_at is None:
        reasons.append("invalid_included_at")

    raw_safe = {key: _json_safe(value) for key, value in raw.items()}
    if reasons:
        return None, {
            "source_row_number": row_number,
            "raw_hash": _hash_payload(raw_safe),
            "reason_codes": reasons,
            "raw_payload": raw_safe,
        }

    canonical = {
        "inn": inn,
        "ogrn": ogrn,
        "registry_number": registry_number,
        "included_at": included_at.isoformat(),
    }
    return {
        "source_row_number": row_number,
        "inn": inn,
        "ogrn": ogrn,
        "registry_number": registry_number,
        "included_at": included_at,
        "row_hash": _hash_payload(canonical),
        "validation_state": "valid",
    }, None


def parse_mintrans_ted_xlsx(
    path: str | Path,
    *,
    source_as_of: datetime,
) -> ParsedSnapshot[MintransTedParsedArtifact]:
    artifact = Path(path)
    if artifact.suffix.lower() != ".xlsx" or not is_zipfile(artifact):
        raise MintransTedSchemaError("Ожидался валидный XLSX artifact")
    source_as_of = _aware_utc(source_as_of)
    try:
        workbook = load_workbook(artifact, read_only=True, data_only=True)
    except Exception as error:
        raise MintransTedSchemaError(f"Не удалось открыть XLSX: {error}") from error

    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header = tuple(next(rows))
        except StopIteration as error:
            raise MintransTedSchemaError("XLSX не содержит строки заголовка") from error
        columns = _resolve_columns(header)
        entries = []
        quarantine = []
        seen_hashes = set()
        duplicates = 0
        source_records = 0
        for row_number, values in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            source_records += 1
            raw = {
                canonical: values[index] if index < len(values) else None
                for canonical, index in columns.items()
            }
            normalized, invalid = _normalize_row(raw, row_number=row_number)
            if invalid is not None:
                quarantine.append(invalid)
                continue
            if normalized["row_hash"] in seen_hashes:
                duplicates += 1
                continue
            seen_hashes.add(normalized["row_hash"])
            entries.append(normalized)
    finally:
        workbook.close()

    coverage = {
        "mode": "fixture",
        "live_ingestion": False,
        "source_records": source_records,
        "normalized_records": len(entries),
        "quarantined_records": len(quarantine),
        "duplicate_records": duplicates,
    }
    payload = MintransTedParsedArtifact(
        entries=tuple(entries),
        quarantine=tuple(quarantine),
        coverage=coverage,
    )
    checksum = validate_artifact(artifact).checksum
    return ParsedSnapshot(
        payload=payload,
        source_as_of=source_as_of,
        record_count=len(entries),
        records_rejected=len(quarantine),
        duplicates=duplicates,
        conflicts=0,
        coverage=coverage,
        version=checksum,
    )


def resolve_exact_identity(
    *,
    inn: str | None,
    ogrn: str | None,
    companies_by_inn: Mapping[str, int],
    companies_by_ogrn: Mapping[str, int],
) -> IdentityMatch:
    inn_company = companies_by_inn.get(inn) if inn else None
    ogrn_company = companies_by_ogrn.get(ogrn) if ogrn else None
    if inn_company is not None and ogrn_company is not None and inn_company != ogrn_company:
        return IdentityMatch(None, "conflict_identity", None)
    if inn_company is not None:
        return IdentityMatch(inn_company, "matched", "inn_exact")
    if ogrn_company is not None:
        return IdentityMatch(ogrn_company, "matched", "ogrn_exact")
    return IdentityMatch(None, "unmatched", None)


def _identity_maps(session, entries: tuple[dict, ...]) -> tuple[dict, dict]:
    inns = {entry["inn"] for entry in entries if entry["inn"]}
    ogrns = {entry["ogrn"] for entry in entries if entry["ogrn"]}
    by_inn = {}
    by_ogrn = {}
    if inns:
        by_inn = dict(
            session.execute(select(Company.inn, Company.id).where(Company.inn.in_(inns))).all()
        )
    if ogrns:
        by_ogrn = dict(
            session.execute(
                select(Company.ogrn, Company.id).where(Company.ogrn.in_(ogrns))
            ).all()
        )
    return by_inn, by_ogrn


def build_registry_listing_fact(
    entry: Mapping[str, object],
    *,
    match_method: str,
    artifact_sha256: str,
) -> tuple[dict, dict]:
    value = {
        "listed": True,
        "registry_number": entry["registry_number"],
        "included_at": entry["included_at"].isoformat(),
    }
    evidence = {
        "source_code": DATASET_CODE,
        "artifact_sha256": artifact_sha256,
        "source_row_hash": entry["row_hash"],
        "matching_method": match_method,
        "source_inn": entry["inn"],
        "source_ogrn": entry["ogrn"],
        "negative_inference": False,
    }
    return value, evidence


def _artifact_row(session, dataset_id: int, staged: StagedFixtureArtifact):
    artifact = session.scalar(
        select(MintransTedRawArtifact).where(
            MintransTedRawArtifact.dataset_id == dataset_id,
            MintransTedRawArtifact.sha256 == staged.checksum,
        )
    )
    if artifact is None:
        artifact = MintransTedRawArtifact(
            dataset_id=dataset_id,
            sha256=staged.checksum,
            original_file_name=staged.manifest["original_file_name"],
            stored_path=str(staged.path),
            media_type=MEDIA_TYPE,
            size_bytes=staged.size_bytes,
            manifest=staged.manifest,
            source_metadata=staged.source_metadata,
        )
        session.add(artifact)
        session.flush()
        return artifact
    expected = {
        "original_file_name": staged.manifest["original_file_name"],
        "stored_path": str(staged.path),
        "media_type": MEDIA_TYPE,
        "size_bytes": staged.size_bytes,
        "manifest": staged.manifest,
        "source_metadata": staged.source_metadata,
    }
    if any(getattr(artifact, key) != value for key, value in expected.items()):
        raise ArtifactImmutabilityError("Метаданные существующего RAW artifact изменились")
    return artifact


def publish_mintrans_ted_snapshot(
    session,
    parsed: MintransTedParsedArtifact,
    *,
    staged: StagedFixtureArtifact,
) -> dict:
    """Replace current normalized rows and facts within the caller transaction."""

    dataset = session.scalar(select(DataSet).where(DataSet.code == DATASET_CODE))
    if dataset is None:
        raise ValueError(f"Dataset не найден: {DATASET_CODE}")
    artifact = _artifact_row(session, dataset.id, staged)
    by_inn, by_ogrn = _identity_maps(session, parsed.entries)

    session.execute(
        delete(TransportForwardingRegistryListing).where(
            TransportForwardingRegistryListing.dataset_id == dataset.id
        )
    )
    session.execute(delete(MintransTedEntry).where(MintransTedEntry.dataset_id == dataset.id))
    session.execute(
        delete(MintransTedQuarantineRow).where(
            MintransTedQuarantineRow.artifact_id == artifact.id
        )
    )

    for invalid in parsed.quarantine:
        session.add(
            MintransTedQuarantineRow(
                dataset_id=dataset.id,
                artifact_id=artifact.id,
                **invalid,
            )
        )

    matched = 0
    unmatched = 0
    conflicts = 0
    pending_facts = []
    observed_at = datetime.now(timezone.utc)
    for values in parsed.entries:
        match = resolve_exact_identity(
            inn=values["inn"],
            ogrn=values["ogrn"],
            companies_by_inn=by_inn,
            companies_by_ogrn=by_ogrn,
        )
        entry = MintransTedEntry(
            dataset_id=dataset.id,
            artifact_id=artifact.id,
            company_id=match.company_id,
            match_state=match.state,
            match_method=match.method,
            **values,
        )
        session.add(entry)
        pending_facts.append((entry, values, match))
        if match.state == "matched":
            matched += 1
        elif match.state == "conflict_identity":
            conflicts += 1
        else:
            unmatched += 1
    session.flush()

    for entry, values, match in pending_facts:
        if match.state != "matched":
            continue
        value, evidence = build_registry_listing_fact(
            values,
            match_method=match.method,
            artifact_sha256=staged.checksum,
        )
        session.add(
            TransportForwardingRegistryListing(
                company_id=match.company_id,
                dataset_id=dataset.id,
                source_entry_id=entry.id,
                fact_code=MINTRANS_TED_FACT_CODE,
                value=value,
                evidence=evidence,
                effective_from=values["included_at"],
                observed_at=observed_at,
            )
        )

    coverage = {
        **parsed.coverage,
        "matched_records": matched,
        "unmatched_records": unmatched,
        "identity_conflicts": conflicts,
        "projected_facts": matched,
    }
    parsed.coverage.update(coverage)
    publication = session.scalar(
        select(DatasetPublication)
        .where(
            DatasetPublication.dataset_id == dataset.id,
            DatasetPublication.status == "staging",
        )
        .order_by(DatasetPublication.id.desc())
        .limit(1)
    )
    if publication is not None:
        metadata = dict(publication.metadata_json or {})
        metadata["coverage"] = coverage
        metadata["conflicts"] = conflicts
        publication.metadata_json = metadata
    return coverage


def run_mintrans_ted_fixture(
    source_path: str | Path,
    *,
    artifact_store: str | Path,
    source_as_of: datetime,
    source_metadata: Mapping[str, object] | None = None,
    expected_sha256: str | None = None,
) -> int:
    """Run only a caller-supplied local fixture; no downloader is reachable here."""

    staged = stage_fixture_artifact(
        source_path,
        artifact_store=artifact_store,
        source_as_of=source_as_of,
        source_metadata=source_metadata,
        expected_sha256=expected_sha256,
    )
    pipeline = BulkUpdatePipeline(
        download=lambda: staged.path,
        parse=lambda path: parse_mintrans_ted_xlsx(path, source_as_of=source_as_of),
        publish=lambda session, payload: publish_mintrans_ted_snapshot(
            session, payload, staged=staged
        ),
        expected_sha256=staged.checksum,
    )
    return pipeline.run(DATASET_CODE, trigger="fixture")
