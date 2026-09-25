from __future__ import annotations

from datetime import date, datetime, timezone
import hashlib
import io
from itertools import chain
import json
import re
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.dialects.postgresql import insert

from app.database.postgres import get_session
from app.models.roskomnadzor import RoskomnadzorCompanyFact, RoskomnadzorPdOperatorCheck, RoskomnadzorPrivatePersonRecord
from app.models.source import DataSet, IngestionRun
from app.services.roskomnadzor_registry_service import DATASETS

_INN = re.compile(r"\d{10}|\d{12}")
_CONTACT_WORDS = ("email", "e-mail", "phone", "телефон", "contact", "контакт", "fio", "фио")


def _local(tag):
    return tag.rsplit("}", 1)[-1].strip().lower()


def _date(value):
    value = str(value or "").strip()
    for pattern in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(value[:10], pattern).date()
        except ValueError:
            pass
    return None


def _flat(element):
    result = {}
    for child in element.iter():
        if child is element or list(child):
            continue
        key, value = _local(child.tag), " ".join(str(child.text or "").split())
        if not value:
            continue
        old = result.get(key)
        result[key] = value if old is None else ([old, value] if not isinstance(old, list) else old + [value])
    return result


def _values(raw, names):
    out = []
    for name in names:
        value = raw.get(name)
        out.extend(value if isinstance(value, list) else ([value] if value else []))
    return [str(item).strip() for item in out if str(item).strip()]


def _first(raw, names):
    values = _values(raw, names)
    return values[0] if values else None


def _safe_details(raw):
    allowed = ("service", "услуг", "territ", "террит", "domain", "домен", "media", "сми", "type", "вид", "form", "форм")
    return {key: value for key, value in raw.items() if any(token in key for token in allowed) and not any(token in key for token in _CONTACT_WORDS)}


def _record_key(channel, inn, number, raw):
    return hashlib.sha256(json.dumps([channel, inn, number, raw], ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def _classify(raw, *, channel, data_date, ordinal):
    canonical = dict(raw)
    for key, value in raw.items():
        compact = re.sub(r"[^a-zа-я0-9]", "", key.lower())
        if compact.endswith("inn") or compact.endswith("инн"):
            canonical.setdefault("inn", value)
        elif compact.endswith("ogrn") or compact.endswith("огрн") or compact.endswith("ogrnip"):
            canonical.setdefault("ogrn", value)
        elif compact in {"entrynum", "registrynumber", "registrationnumber", "licensenumber", "licensenum", "licencenum", "regnum"}:
            canonical.setdefault("registry_number", value)
        elif compact in {"licstatusname", "statusname"}:
            canonical.setdefault("status", value)
        elif compact.endswith("name") and "person" not in compact and not any(token in compact for token in ("service", "territory", "media", "smi")):
            canonical.setdefault("organization_name", value)
        elif compact.endswith("date") and "end" not in compact:
            canonical.setdefault("registration_date", value)
        elif compact in {"datestart", "issuedate"}:
            canonical.setdefault("registration_date", value)
        elif compact in {"dateend", "validuntil", "expirydate"}:
            canonical.setdefault("valid_until", value)
    raw = canonical
    inns = _values(raw, ("inn", "инн", "founder_inn", "inn_founder", "uch_inn"))
    inn = next((value for value in inns if _INN.fullmatch(value)), None)
    ogrn = _first(raw, ("ogrn", "огрн", "ogrnip", "огрнип"))
    number = _first(raw, ("license_number", "licensenumber", "number", "reg_number", "registration_number", "registry_number", "regnum", "свидетельство")) or f"row-{ordinal}"
    name = _first(raw, ("name", "full_name", "organization_name", "licensee", "founder", "наименование"))
    key = _record_key(channel, inn, number, raw)
    if inn and len(inn) == 10 and (not ogrn or len(ogrn) in {0, 13}):
        return "public", {
            "data_date": data_date, "record_key": key, "channel": channel,
            "inn": inn, "ogrn": ogrn or None, "external_number": number,
            "name": name, "status": _first(raw, ("status", "state", "статус")),
            "issued_at": _date(_first(raw, ("date", "issue_date", "registration_date", "дата"))),
            "valid_until": _date(_first(raw, ("date_end", "valid_until", "expiry_date"))),
            "public_details": _safe_details(raw),
        }
    identifier = inn or ogrn or f"unidentified:{ordinal}"
    return "private", {
        "data_date": data_date, "record_key": key, "channel": channel,
        "identifier_kind": "inn12_or_ogrnip" if inn or ogrn else "unidentified",
        "identifier_hash": hashlib.sha256(identifier.encode()).hexdigest(),
        "private_payload": raw, "is_published": False,
    }


def parse_xml_snapshot(content: bytes, *, channel: str, data_date: date, record_tags: set[str] | None = None):
    public, private, seen = [], [], set()
    duplicates = rejected = source = 0
    tags = {item.lower() for item in (record_tags or {"license", "licenses", "record", "item", "informationdistributor", "massmedia", "media"})}
    try:
        for _, element in ET.iterparse(io.BytesIO(content), events=("end",)):
            if _local(element.tag) not in tags:
                continue
            raw = _flat(element)
            if not raw:
                element.clear()
                continue
            candidates = [raw]
            if channel == "media":
                media_name = next((str(child.text or "").strip() for child in element if _local(child.tag) == "name"), None)
                founders = [child for child in element.iter() if _local(child.tag) == "founder"]
                candidates = []
                for founder in founders or [None]:
                    founder_raw = _flat(founder) if founder is not None else {}
                    candidates.append({
                        **raw,
                        "inn": _first(founder_raw, ("inn", "инн")),
                        "organization_name": _first(founder_raw, ("name", "наименование")),
                        "media_name": media_name,
                    })
            for candidate in candidates:
                source += 1
                kind, row = _classify(candidate, channel=channel, data_date=data_date, ordinal=source)
                if row["record_key"] in seen:
                    duplicates += 1
                else:
                    seen.add(row["record_key"])
                    (public if kind == "public" else private).append(row)
            element.clear()
    except ET.ParseError as error:
        raise ValueError("Некорректный или неполный XML Роскомнадзора") from error
    if not source:
        raise ValueError("В XML Роскомнадзора не найдены записи ожидаемой структуры")
    return {"public_records": public, "private_records": private, "source_records": source, "imported_records": len(public) + len(private), "duplicate_records": duplicates, "rejected_records": rejected}


def parse_hosting_xlsx(content: bytes, *, data_date: date):
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("Некорректный XLSX реестра провайдеров хостинга") from error
    sheet = workbook.active
    row_iterator = iter(sheet.iter_rows(values_only=True))
    prefix = []
    for _ in range(20):
        try:
            prefix.append(next(row_iterator))
        except StopIteration:
            break
    header_index = next((index for index, row in enumerate(prefix) if any("инн" in str(value or "").lower() for value in row)), None)
    if header_index is None:
        raise ValueError("В XLSX не найден заголовок с ИНН")
    headers = [" ".join(str(value or "").lower().split()) for value in prefix[header_index]]
    public, private, seen = [], [], set()
    source = duplicates = 0
    for values in chain(prefix[header_index + 1:], row_iterator):
        if not any(value not in (None, "") for value in values):
            continue
        source += 1
        raw = {headers[i] or f"column_{i}": str(value).strip() for i, value in enumerate(values) if value not in (None, "")}
        # Canonical aliases keep the classifier independent from column wording.
        for key, value in list(raw.items()):
            if "инн" in key: raw.setdefault("inn", re.sub(r"\D", "", value))
            elif "огрн" in key: raw.setdefault("ogrn", re.sub(r"\D", "", value))
            elif "реестр" in key or "номер" in key: raw.setdefault("registry_number", value)
            elif "наимен" in key: raw.setdefault("organization_name", value)
        kind, row = _classify(raw, channel="hosting", data_date=data_date, ordinal=source)
        if row["record_key"] in seen: duplicates += 1
        else:
            seen.add(row["record_key"]); (public if kind == "public" else private).append(row)
    workbook.close()
    return {"public_records": public, "private_records": private, "source_records": source, "imported_records": len(public) + len(private), "duplicate_records": duplicates, "rejected_records": 0}


def validate_complete_snapshot(parsed):
    if not parsed["source_records"] or not parsed["imported_records"]:
        raise ValueError("Пустой snapshot Роскомнадзора не публикуется")
    if parsed["rejected_records"]:
        raise ValueError("Snapshot содержит отклонённые записи")
    if parsed["source_records"] != parsed["imported_records"] + parsed["duplicate_records"]:
        raise ValueError("Контроль количества snapshot не пройден")
    if any(row.get("is_published") for row in parsed["private_records"]):
        raise ValueError("Закрытая person-запись не может быть опубликована")


def publish_snapshot(*, dataset_code, parsed, data_date, checksum=None, source_file_name=None, details=None, batch_size=1000):
    validate_complete_snapshot(parsed)
    session = get_session()
    try:
        dataset = session.scalar(select(DataSet).where(DataSet.code == dataset_code).with_for_update())
        if dataset is None: raise ValueError(f"Dataset {dataset_code} не зарегистрирован")
        session.execute(delete(RoskomnadzorCompanyFact).where(RoskomnadzorCompanyFact.dataset_id == dataset.id))
        session.execute(delete(RoskomnadzorPrivatePersonRecord).where(RoskomnadzorPrivatePersonRecord.dataset_id == dataset.id))
        for model, records in ((RoskomnadzorCompanyFact, parsed["public_records"]), (RoskomnadzorPrivatePersonRecord, parsed["private_records"])):
            for start in range(0, len(records), batch_size):
                session.execute(insert(model), [{"dataset_id": dataset.id, **row} for row in records[start:start + batch_size]])
        now = datetime.now(timezone.utc)
        session.add(IngestionRun(dataset_id=dataset.id, status="success", started_at=now, finished_at=now, data_date=data_date, source_file_name=source_file_name, source_url=dataset.source_url, file_checksum=checksum, rows_read=parsed["source_records"], rows_inserted=parsed["imported_records"], rows_updated=0, rows_skipped=parsed["duplicate_records"], errors_count=0, details={"public_records": len(parsed["public_records"]), "private_records": len(parsed["private_records"]), "privacy_boundary": "separate_unpublished_table", "complete_snapshot": True, **(details or {})}))
        dataset.last_data_date, dataset.last_success_at = data_date, now
        session.commit()
        return {"dataset_id": dataset.id, "public": len(parsed["public_records"]), "private": len(parsed["private_records"])}
    except Exception:
        session.rollback(); raise
    finally:
        session.close()


def save_pd_operator_check(*, inn, request_date, **values):
    session = get_session()
    try:
        dataset_id = session.scalar(select(DataSet.id).where(DataSet.code == DATASETS["pd_operators"]))
        now = datetime.now(timezone.utc)
        all_values = {"dataset_id": dataset_id, "inn": inn, "request_date": request_date, **values, "checked_at": now, "updated_at": now}
        session.execute(insert(RoskomnadzorPdOperatorCheck).values(**all_values).on_conflict_do_update(constraint="uq_rkn_pd_operator_check", set_={k: v for k, v in all_values.items() if k not in {"dataset_id", "inn", "request_date"}}))
        session.commit()
        return session.scalar(select(RoskomnadzorPdOperatorCheck).where(RoskomnadzorPdOperatorCheck.dataset_id == dataset_id, RoskomnadzorPdOperatorCheck.inn == inn, RoskomnadzorPdOperatorCheck.request_date == request_date))
    except Exception:
        session.rollback(); raise
    finally:
        session.close()
