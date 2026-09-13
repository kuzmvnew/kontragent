import argparse
import csv
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.dialects.postgresql import insert

from app.database.postgres import get_session
from app.models.company import (
    Company,
    CompanyBranch,
    CompanyContact,
    CompanyFinancial,
    CompanyIdentifier,
    CompanyManager,
)


SOURCE_NAME = "excel_import"


BRANCH_WORDS = (
    "филиал",
    "подразделение",
    "точка продаж",
    "магазин",
    "офис",
    "обособленное подразделение",
)


def clean_text(value):
    if value is None:
        return None

    value = str(value).strip()

    if not value:
        return None

    return value


def clean_digits(value):
    value = clean_text(value)

    if value is None:
        return None

    value = re.sub(r"\D", "", value)

    return value or None


def clean_integer(value):
    value = clean_text(value)

    if value is None:
        return None

    value = value.replace(" ", "")
    value = value.replace("\xa0", "")
    value = value.replace(",", ".")

    try:
        return int(float(value))
    except (ValueError, TypeError):
        return None


def split_values(value):
    value = clean_text(value)

    if not value:
        return []

    reader = csv.reader(
        [value],
        skipinitialspace=True,
    )

    result = next(reader)

    cleaned = []

    for item in result:
        item = clean_text(item)

        if item:
            cleaned.append(item)

    return cleaned


def make_full_name(last_name, first_name, middle_name):
    parts = [
        clean_text(last_name),
        clean_text(first_name),
        clean_text(middle_name),
    ]

    parts = [
        part
        for part in parts
        if part
    ]

    if not parts:
        return None

    return " ".join(parts)


def looks_like_branch(row):
    branch_code = clean_text(
        row.get("Код филиала")
    )

    if branch_code:
        return True

    name = (
        clean_text(
            row.get("Название")
        )
        or ""
    ).lower()

    return any(
        word in name
        for word in BRANCH_WORDS
    )


def company_row_score(row):
    """
    Определяет, какую строку считать основной
    для одного ИНН.

    Чем выше результат, тем вероятнее,
    что это головная организация,
    а не филиал или точка продаж.
    """

    score = 0

    if not looks_like_branch(row):
        score += 100

    if clean_text(row.get("Код филиала")) is None:
        score += 50

    if clean_text(row.get("ОКПО")):
        score += 10

    if clean_text(row.get("Фамилия руководителя")):
        score += 10

    if clean_text(row.get("Выручка")):
        score += 5

    if clean_text(row.get("Стоимость")):
        score += 5

    if clean_text(row.get("Сайт")):
        score += 3

    return score


def read_excel(file_path):
    print()
    print("Читаем Excel...")

    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    rows = worksheet.iter_rows(
        values_only=True
    )

    headers = next(rows)

    headers = [
        clean_text(header)
        for header in headers
    ]

    grouped = defaultdict(list)

    total_rows = 0
    skipped_rows = 0

    for values in rows:
        total_rows += 1

        row = dict(
            zip(headers, values)
        )

        inn = clean_digits(
            row.get("ИНН")
        )

        name = clean_text(
            row.get("Название")
        )

        if (
            not inn
            or len(inn) not in (10, 12)
            or not name
        ):
            skipped_rows += 1
            continue

        row["ИНН"] = inn

        grouped[inn].append(row)

    workbook.close()

    print(
        f"Строк прочитано: {total_rows:,}"
        .replace(",", " ")
    )

    print(
        f"Уникальных ИНН: {len(grouped):,}"
        .replace(",", " ")
    )

    print(
        f"Некорректных строк: {skipped_rows:,}"
        .replace(",", " ")
    )

    return grouped


def upsert_company(session, row):
    inn = clean_digits(
        row.get("ИНН")
    )

    company_data = {
        "inn": inn,
        "kpp": clean_digits(
            row.get("КПП")
        ),
        "ogrn": clean_digits(
            row.get("ОГРН")
        ),
        "okpo": clean_digits(
            row.get("ОКПО")
        ),
        "name": clean_text(
            row.get("Название")
        ),
        "address": clean_text(
            row.get("Адрес")
        ),
        "activity": clean_text(
            row.get("Вид деятельности")
        ),
        "website": clean_text(
            row.get("Сайт")
        ),
        "source": SOURCE_NAME,
    }

    statement = insert(
        Company
    ).values(
        **company_data
    )

    statement = statement.on_conflict_do_update(
        index_elements=[
            Company.inn
        ],
        set_={
            "kpp": company_data["kpp"],
            "ogrn": company_data["ogrn"],
            "okpo": company_data["okpo"],
            "name": company_data["name"],
            "address": company_data["address"],
            "activity": company_data["activity"],
            "website": company_data["website"],
            "source": SOURCE_NAME,
        },
    )

    statement = statement.returning(
        Company.id
    )

    company_id = session.execute(
        statement
    ).scalar_one()

    return company_id


def clear_excel_children(
    session,
    company_id,
):
    session.execute(
        delete(
            CompanyManager
        ).where(
            CompanyManager.company_id
            == company_id,
            CompanyManager.source
            == SOURCE_NAME,
        )
    )

    session.execute(
        delete(
            CompanyContact
        ).where(
            CompanyContact.company_id
            == company_id,
            CompanyContact.source
            == SOURCE_NAME,
        )
    )

    session.execute(
        delete(
            CompanyFinancial
        ).where(
            CompanyFinancial.company_id
            == company_id,
            CompanyFinancial.source
            == SOURCE_NAME,
        )
    )

    session.execute(
        delete(
            CompanyIdentifier
        ).where(
            CompanyIdentifier.company_id
            == company_id,
            CompanyIdentifier.source
            == SOURCE_NAME,
        )
    )

    session.execute(
        delete(
            CompanyBranch
        ).where(
            CompanyBranch.company_id
            == company_id,
            CompanyBranch.source
            == SOURCE_NAME,
        )
    )


def add_manager(
    session,
    company_id,
    row,
):
    last_name = clean_text(
        row.get("Фамилия руководителя")
    )

    first_name = clean_text(
        row.get("Имя руководителя")
    )

    middle_name = clean_text(
        row.get("Отчество руководителя")
    )

    full_name = make_full_name(
        last_name,
        first_name,
        middle_name,
    )

    if not full_name:
        return

    session.add(
        CompanyManager(
            company_id=company_id,
            last_name=last_name,
            first_name=first_name,
            middle_name=middle_name,
            full_name=full_name,
            is_current=True,
            source=SOURCE_NAME,
        )
    )


def add_financial(
    session,
    company_id,
    row,
):
    revenue = clean_integer(
        row.get("Выручка")
    )

    company_value = clean_integer(
        row.get("Стоимость")
    )

    employee_count = clean_integer(
        row.get("Количество сотрудников")
    )

    if (
        revenue is None
        and company_value is None
        and employee_count is None
    ):
        return

    session.add(
        CompanyFinancial(
            company_id=company_id,
            year=None,
            revenue=revenue,
            company_value=company_value,
            employee_count=employee_count,
            source=SOURCE_NAME,
        )
    )


def add_contacts(
    session,
    company_id,
    rows,
):
    contacts = set()

    for row in rows:

        phones = split_values(
            row.get("Телефоны")
        )

        emails = split_values(
            row.get("email")
        )

        websites = split_values(
            row.get("Сайт")
        )

        for phone in phones:
            contacts.add(
                ("phone", phone)
            )

        for email in emails:
            contacts.add(
                (
                    "email",
                    email.lower(),
                )
            )

        for website in websites:
            contacts.add(
                ("website", website)
            )

    sorted_contacts = sorted(
        contacts
    )

    first_types = set()

    for contact_type, value in sorted_contacts:

        is_primary = (
            contact_type
            not in first_types
        )

        if is_primary:
            first_types.add(
                contact_type
            )

        session.add(
            CompanyContact(
                company_id=company_id,
                contact_type=contact_type,
                value=value,
                is_primary=is_primary,
                source=SOURCE_NAME,
            )
        )


def add_identifier(
    session,
    company_id,
    identifier_type,
    value,
):
    value = clean_text(value)

    if not value:
        return

    session.add(
        CompanyIdentifier(
            company_id=company_id,
            identifier_type=identifier_type,
            value=value,
            source=SOURCE_NAME,
        )
    )


def add_identifiers(
    session,
    company_id,
    rows,
):
    identifiers = set()

    for row in rows:

        edo_values = split_values(
            row.get("Идентификатор ЭДО")
        )

        for value in edo_values:
            identifiers.add(
                ("edo", value)
            )

        pfr = clean_text(
            row.get("Рег. номер ПФ")
        )

        if pfr:
            identifiers.add(
                ("pfr", pfr)
            )

        egais = clean_text(
            row.get("ЕГАИС")
        )

        if egais:
            identifiers.add(
                ("egais", egais)
            )

        gln = clean_text(
            row.get("GLN")
        )

        if gln:
            identifiers.add(
                ("gln", gln)
            )

    for identifier_type, value in sorted(
        identifiers
    ):
        add_identifier(
            session,
            company_id,
            identifier_type,
            value,
        )


def add_branches(
    session,
    company_id,
    rows,
    canonical_row,
):
    for row in rows:

        if row is canonical_row:
            continue

        if not looks_like_branch(row):
            continue

        session.add(
            CompanyBranch(
                company_id=company_id,
                branch_code=clean_text(
                    row.get("Код филиала")
                ),
                name=clean_text(
                    row.get("Название")
                ),
                address=clean_text(
                    row.get("Адрес")
                ),
                kpp=clean_digits(
                    row.get("КПП")
                ),
                source=SOURCE_NAME,
            )
        )

    if (
        looks_like_branch(canonical_row)
        and clean_text(
            canonical_row.get(
                "Код филиала"
            )
        )
    ):
        session.add(
            CompanyBranch(
                company_id=company_id,
                branch_code=clean_text(
                    canonical_row.get(
                        "Код филиала"
                    )
                ),
                name=clean_text(
                    canonical_row.get(
                        "Название"
                    )
                ),
                address=clean_text(
                    canonical_row.get(
                        "Адрес"
                    )
                ),
                kpp=clean_digits(
                    canonical_row.get(
                        "КПП"
                    )
                ),
                source=SOURCE_NAME,
            )
        )


def import_companies(file_path):
    grouped = read_excel(
        file_path
    )

    session = get_session()

    imported = 0
    failed = 0

    print()
    print("Начинаем импорт в PostgreSQL...")
    print()

    try:
        for inn, rows in grouped.items():

            try:
                canonical_row = max(
                    rows,
                    key=company_row_score,
                )

                company_id = upsert_company(
                    session,
                    canonical_row,
                )

                clear_excel_children(
                    session,
                    company_id,
                )

                add_manager(
                    session,
                    company_id,
                    canonical_row,
                )

                add_financial(
                    session,
                    company_id,
                    canonical_row,
                )

                add_contacts(
                    session,
                    company_id,
                    rows,
                )

                add_identifiers(
                    session,
                    company_id,
                    rows,
                )

                add_branches(
                    session,
                    company_id,
                    rows,
                    canonical_row,
                )

                session.commit()

                imported += 1

                if imported % 250 == 0:
                    print(
                        f"Импортировано: "
                        f"{imported:,} компаний"
                        .replace(",", " ")
                    )

            except Exception as error:
                session.rollback()

                failed += 1

                print()
                print(
                    f"ОШИБКА для ИНН {inn}:"
                )
                print(error)
                print()

        print()
        print("======================================")
        print("ИМПОРТ ЗАВЕРШЁН")
        print("======================================")

        print(
            "Компаний импортировано:",
            f"{imported:,}".replace(
                ",",
                " ",
            ),
        )

        print(
            "Ошибок:",
            f"{failed:,}".replace(
                ",",
                " ",
            ),
        )

        print()

    finally:
        session.close()


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Импорт компаний из Excel "
            "в PostgreSQL"
        )
    )

    parser.add_argument(
        "file",
        help="Путь к Excel-файлу",
    )

    args = parser.parse_args()

    file_path = Path(
        args.file
    ).expanduser().resolve()

    if not file_path.exists():
        raise FileNotFoundError(
            f"Файл не найден: {file_path}"
        )

    print(
        f"Файл: {file_path}"
    )

    import_companies(
        file_path
    )


if __name__ == "__main__":
    main()